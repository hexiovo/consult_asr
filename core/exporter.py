from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path
from typing import Any, Mapping

from docx import Document

from core.transcript import build_sentences, format_sentence, revise_sentences


_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _safe_stem(value: str) -> str:
    stem = _INVALID_FILENAME.sub("_", value).strip(" .")
    if not stem:
        stem = "咨询记录"
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem = f"_{stem}"
    return stem[:120]


def _available_stem(directory: Path, requested: str) -> str:
    base = _safe_stem(requested)
    candidate = base
    suffix = 2
    while (directory / f"{candidate}.docx").exists():
        candidate = f"{base} ({suffix})"
        suffix += 1
    return candidate


def _temporary_path(directory: Path, suffix: str) -> Path:
    with tempfile.NamedTemporaryFile(
        prefix=".consult-asr-", suffix=suffix, dir=directory, delete=False
    ) as file:
        return Path(file.name)


def _stamp(seconds: float) -> str:
    whole = max(0, int(seconds))
    return f"{whole // 60:02d}:{whole % 60:02d}"


def merge_adjacent_turns(words: list[dict[str, Any]], max_gap_seconds: float = 2.0) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    for word in sorted(words, key=lambda item: (float(item["start"]), float(item["end"]))):
        item = {
            "start": float(word["start"]),
            "end": float(word["end"]),
            "text": str(word.get("text", "")),
            "speaker_id": word.get("speaker_id"),
            "uncertain": bool(word.get("uncertain", False)),
        }
        if (
            merged
            and merged[-1]["speaker_id"] == item["speaker_id"]
            and item["start"] - merged[-1]["end"] <= max_gap_seconds
        ):
            merged[-1]["end"] = max(merged[-1]["end"], item["end"])
            merged[-1]["text"] += item["text"]
            merged[-1]["uncertain"] = merged[-1]["uncertain"] or item["uncertain"]
        else:
            merged.append(item)
    return merged


def write_exports(
    *,
    words: list[dict[str, Any]],
    output_dir: str | os.PathLike[str],
    filename: str,
    source_name: str,
    duration_seconds: float | None,
    status: str,
    role_labels: Mapping[str, str] | None = None,
    roles_confirmed: bool = False,
    segments: list[dict[str, Any]] | None = None,
    emotion_windows: list[dict[str, Any]] | None = None,
    analysis_windows: list[dict[str, Any]] | None = None,
    tone_enabled: bool | None = None,
    tone_status: str | None = None,
    tone_error: str = "",
    tone_model: Mapping[str, Any] | None = None,
    sentence_revisions: Mapping[str, Mapping[str, Any]] | None = None,
    revision: int = 0,
    target_path: str | os.PathLike[str] | None = None,
    overwrite: bool = False,
) -> Path:
    """Write only a Word transcript, optionally to an explicitly selected path."""
    directory = Path(output_dir).expanduser().resolve()
    if isinstance(revision, bool) or not isinstance(revision, int) or revision < 0:
        raise ValueError("revision must be a non-negative integer")
    if target_path is None:
        directory.mkdir(parents=True, exist_ok=True)
        stem = _available_stem(directory, filename)
        word_path = directory / f"{stem}.docx"
    else:
        word_path = Path(target_path).expanduser().resolve(strict=False)
        if word_path.parent != directory or word_path.suffix.lower() != ".docx":
            raise ValueError("Word target must be a .docx file inside output_dir")
        word_path.parent.mkdir(parents=True, exist_ok=True)
    temp_word: Path | None = None

    speaker_names = {"speaker_1": "说话人1", "speaker_2": "说话人2"}
    if role_labels:
        speaker_names.update({str(key): str(value) for key, value in role_labels.items() if value})
    turns = merge_adjacent_turns(words) if segments is None else [
        {
            "start": float(turn["start"]),
            "end": float(turn["end"]),
            "text": str(turn.get("text", "")),
            "speaker_id": turn.get("speaker_id"),
            "uncertain": bool(turn.get("uncertain", False)),
            **({"acoustics": turn["acoustics"]} if isinstance(turn.get("acoustics"), dict) else {}),
        }
        for turn in sorted(segments, key=lambda item: (float(item["start"]), float(item["end"])))
    ]
    feature_windows = analysis_windows if analysis_windows is not None else [
        {**turn, "analysis_safe": turn.get("analysis_safe", True)}
        for turn in turns if isinstance(turn.get("acoustics"), dict)
    ]
    use_revisions = sentence_revisions is not None or revision != 0
    sentences = (
        revise_sentences(words, sentence_revisions, emotion_windows, tone_enabled, feature_windows)
        if use_revisions
        else build_sentences(words, emotion_windows, tone_enabled, feature_windows)
    )
    doc = Document()
    doc.add_heading("咨询记录", level=1)
    doc.add_paragraph(f"音频：{Path(source_name).name}")
    if duration_seconds is not None:
        doc.add_paragraph(f"时长：{_stamp(duration_seconds)}")
    if not turns:
        doc.add_paragraph("未检测到语音。")
    for sentence in sentences:
        speaker_id = sentence["speaker_id"]
        speaker = speaker_names.get(speaker_id, "说话人（待确认）")
        marker = " [待核对]" if sentence["uncertain"] else ""
        timestamp = f"[{_stamp(sentence['start'])}–{_stamp(sentence['end'])}] "
        doc.add_paragraph(f"{timestamp}{speaker}{marker}：{format_sentence(sentence)}")

    has_extended_analysis = any(
        value is not None
        for value in (segments, emotion_windows, analysis_windows, tone_enabled, tone_status, tone_model)
    ) or use_revisions
    if has_extended_analysis:
        doc.add_heading("说明", level=2)
        doc.add_paragraph(
            "声学语速根据可靠窗口内的 ASR 文字估算；情绪标签表示模型给出的语音倾向，不用于诊断。"
            "模型置信分数为未校准的原始输出，不是咨询场景准确率或概率。句末音高走势只描述声学变化，"
            "不据此判定疑问句。每句括号取同一说话人、与句中文字时间重叠最多的可靠窗口及其原始分数；"
            "句中语音未被可靠情绪窗口完整覆盖时显示无法判定。音量、音调/基频、语速、停顿及音高走势统一列在句后括号，取同一说话人的最大重叠可靠声学窗口。声学数值是该窗口的统计，不是对整句重新测量；覆盖不足或数值缺失时显示—。"
        )
        if tone_enabled is False:
            doc.add_paragraph("语音情绪分析：已关闭。")
        else:
            if tone_status in {"failed", "partial"}:
                status_text = "部分失败" if tone_status == "partial" else "失败"
                doc.add_paragraph(f"语音情绪分析：{status_text}。" + (tone_error.strip() or "模型运行失败。"))

    try:
        temp_word = _temporary_path(word_path.parent, ".docx")
        doc.save(temp_word)
        if overwrite:
            os.replace(temp_word, word_path)
        elif os.name == "nt":
            # Windows rename refuses existing targets and supports FAT volumes.
            os.rename(temp_word, word_path)
        else:
            # Publish without replacing a file created after the filename check.
            os.link(temp_word, word_path)
        return word_path
    finally:
        if temp_word is not None:
            try:
                temp_word.unlink(missing_ok=True)
            except OSError:
                pass
